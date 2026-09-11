#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
КЕ-2, верхний блок страницы офисных перегородок (шапка).
SEO-версия того блока, что сейчас стоит наверху страницы.

Собирает:
  shapka.html                        - блок для вставки над содержательной частью
  shapka.md                          - текст для вычитки
  GENGLASS_KE-2_shapka_seo.docx      - версия для заказчика

Рамки те же, что у КЕ-2: цифр в дБ нет, цены за м2 нет, Metal-GM не упоминается,
только дефис, бренд заглавными, ссылки только из разрешённого списка.
"""

import os
import re

H1 = "Офисные перегородки на заказ и комплектация объектов"

ALT_H1 = [
    "Офисные перегородки на заказ и комплектация объектов",
    "Стеклянные офисные перегородки: изготовление и монтаж",
    "Офисные перегородки под ключ: от обмера до сдачи объекта",
]

LINKS = {
    "стационарные перегородки": "/staczionarnye-peregorodki-na-zakaz/",
    "алюминиевые перегородки": "/alyuminievye-peregorodki/",
    "акустические перегородки": "/ofisnye-peregorodki/akusticheskie-peregorodki/",
    "дилерам и подрядчикам": "/dileram/",
}
STOP_URLS = ["/montazh-peregorodok/"]

BLOCKS = [
    ("p",
     "GENGLASS проектирует, изготавливает и монтирует стеклянные офисные "
     "перегородки на собственном производстве в Домодедово. Цикл закрыт "
     "целиком: замер, проект, резка и закалка стекла по ГОСТ 30698-2014, "
     "окраска профиля, сборка и монтаж своей бригадой."),
    ("p",
     "Перегородки для офиса решают задачу планировки, а не отделки. Они делят "
     "этаж на зоны и оставляют свет всей площади: переговорная перестаёт быть "
     "комнатой без окна, а опенспейс не превращается в лабиринт коридоров."),
    ("p",
     "Отдельное направление - комплектация объектов от 100 м2. Заказчик "
     "получает не изделие, а закрытый участок работ: обмерный план, "
     "спецификацию, смету, изготовление очередями под график стройки, монтаж "
     "и сдачу с актом. Один договор и один ответственный за геометрию, сроки "
     "и результат."),
    ("p", "Что входит в работу:"),
    ("ul", [
        "замер на объекте и обмерный план с привязкой к осям и коммуникациям;",
        "схема раскладки, спецификация по позициям и смета;",
        "изготовление на своём производстве, контроль ОТК перед отгрузкой;",
        "доставка, такелаж и подъём на этаж;",
        "монтаж силами своей бригады и сдача зонами.",
    ]),
    ("p",
     "Офисные перегородки на заказ делаем под конкретный проём и высоту этажа. "
     "Глухие [стационарные перегородки](/staczionarnye-peregorodki-na-zakaz/) "
     "закрывают периметр кабинетов, "
     "[алюминиевые перегородки](/alyuminievye-peregorodki/) идут на большом "
     "метраже, где важен лёгкий профиль, а "
     "[акустические перегородки](/ofisnye-peregorodki/akusticheskie-peregorodki/) "
     "ставят там, где в переговорной нужна тишина."),
    ("p",
     "Смету по готовому обмерному плану GENGLASS считает за 1 день. Если плана "
     "нет, сначала выезжает замерщик: без фактических размеров смета остаётся "
     "вилкой, которая на подписании всё равно поедет. Гарантия на конструкцию "
     "и фурнитуру - 12 месяцев, для объектов срок и состав обязательств "
     "фиксируются договором."),
]

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def plain(t):
    return LINK_RE.sub(r"\1", t)


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_html(t):
    return LINK_RE.sub(r'<a href="\2">\1</a>', esc(t))


def parts():
    out = []
    for b in BLOCKS:
        if b[0] == "p":
            out.append(plain(b[1]))
        elif b[0] == "ul":
            out.extend(plain(i) for i in b[1])
    return out


def render_html():
    out = ['<section class="page-intro seo-ofisnye-shapka">',
           f"  <h1>{esc(H1)}</h1>"]
    for b in BLOCKS:
        if b[0] == "p":
            out.append(f"  <p>{to_html(b[1])}</p>")
        elif b[0] == "ul":
            out.append("  <ul>")
            out.extend(f"    <li>{to_html(i)}</li>" for i in b[1])
            out.append("  </ul>")
    out.append("</section>")
    return "\n".join(out) + "\n"


def render_md():
    out = [f"# {H1}", ""]
    for b in BLOCKS:
        if b[0] == "p":
            out += [b[1], ""]
        elif b[0] == "ul":
            out += [f"- {i}" for i in b[1]] + [""]
    return "\n".join(out)


def render_docx(path, volume):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(H1, level=1)
    for b in BLOCKS:
        if b[0] == "p":
            doc.add_paragraph(plain(b[1]))
        elif b[0] == "ul":
            for i in b[1]:
                doc.add_paragraph(plain(i), style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Служебное", level=2)

    doc.add_heading("Варианты H1", level=3)
    for t in ALT_H1:
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("Ссылки в блоке", level=3)
    used = set(a for a, _ in LINK_RE.findall(render_md()))
    for anchor, url in LINKS.items():
        mark = "стоит" if anchor in used else "не используется в шапке"
        doc.add_paragraph(f"{anchor} - {url} ({mark})", style="List Bullet")

    doc.add_heading("Ключевые фразы в блоке", level=3)
    for phrase in KEYS:
        doc.add_paragraph(f"{phrase}: {COUNTS[phrase]}", style="List Bullet")

    note = doc.add_paragraph(
        f"Объём {volume} знаков с пробелами. Блок ставится над содержательной "
        "частью страницы. Рамки КЕ-2 соблюдены: значений в дБ нет, цены за м2 "
        "нет, только дефис, бренд заглавными.")
    note.runs[0].font.size = Pt(9)
    doc.save(path)


KEYS = [
    "офисные перегородки",
    "стеклянные офисные перегородки",
    "перегородки для офиса",
    "офисные перегородки на заказ",
    "комплектация объектов",
]
COUNTS = {}


def selfcheck(text, volume):
    problems, notes = [], []
    notes.append(f"объём: {volume} знаков с пробелами")

    for k in KEYS:
        COUNTS[k] = len(re.findall(k, text, re.I))
        if COUNTS[k] == 0:
            problems.append(f"нет ключа: {k}")
    notes.append("ключи: " + ", ".join(f"{k} - {COUNTS[k]}" for k in KEYS))

    for dash in ("—", "–"):
        if dash in text:
            problems.append(f"найдено тире {dash!r}")
    if re.search(r"\d\s*дб", text, re.I):
        problems.append("значение в дБ")
    if re.search(r"(от|за)\s*\d[\d\s]*\s*(₽|руб)", text, re.I):
        problems.append("в шапке появилась цена")
    for w in ("преимуществ", "единственн", "лучший", "лучшая", "лучшие"):
        if w in text.lower():
            problems.append(f"стоп-слово: {w}")
    if "metal-gm" in text.lower():
        problems.append("упомянут Metal-GM")
    if re.search(r"Genglass|GenGlass", text):
        problems.append("бренд не заглавными")

    md = render_md()
    urls = set(LINK_RE.findall(md))
    for anchor, url in urls:
        if LINKS.get(anchor) != url:
            problems.append(f"недопустимая ссылка: [{anchor}]({url})")
    for s in STOP_URLS:
        if s in md:
            problems.append(f"ссылка из стоп-листа: {s}")
    notes.append(f"ссылок в блоке: {len(urls)} из 4 разрешённых")

    brand = len(re.findall(r"GENGLASS", text))
    notes.append(f"GENGLASS: {brand} раз")
    if "12 месяцев" not in text:
        problems.append("не названа гарантия")
    return notes, problems


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ps = parts()
    volume = sum(len(p) for p in ps) + len(H1)
    text = " ".join([H1] + ps)
    notes, problems = selfcheck(text, volume)

    with open(os.path.join(here, "shapka.html"), "w", encoding="utf-8") as f:
        f.write(render_html())
    with open(os.path.join(here, "shapka.md"), "w", encoding="utf-8") as f:
        f.write(render_md())
    render_docx(os.path.join(here, "GENGLASS_KE-2_shapka_seo.docx"), volume)

    print("\n".join("  " + n for n in notes))
    print()
    print("ПРОБЛЕМЫ:\n" + "\n".join("  - " + p for p in problems)
          if problems else "Самопроверка шапки: без замечаний")


if __name__ == "__main__":
    main()
